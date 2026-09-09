"""Conversion engine for common text, image, PDF, Office, media and archive formats."""
from __future__ import annotations

import csv
import html
import json
import shutil
import subprocess
import tarfile
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path
import xml.etree.ElementTree as ET

import fitz
import yaml
from PIL import Image
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt

from .dependencies import find_soffice

TEXT = {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".xml"}
IMAGES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif", ".gif"}
OFFICE = {".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".odt", ".ods", ".odp", ".rtf"}
MEDIA = {".mp3", ".wav", ".flac", ".aac", ".ogg", ".m4a", ".opus", ".mp4", ".mkv", ".avi", ".mov", ".webm", ".mpeg", ".mpg", ".m4v"}
ARCHIVES = {".zip", ".tar", ".gz", ".bz2", ".xz", ".tgz"}
ALL_FORMATS = sorted(TEXT | IMAGES | OFFICE | MEDIA | ARCHIVES | {".pdf"})


def format_name(path: Path) -> str:
    name = path.name.lower()
    if name.endswith(".tar.gz"): return ".tar.gz"
    if name.endswith(".tar.bz2"): return ".tar.bz2"
    if name.endswith(".tar.xz"): return ".tar.xz"
    return path.suffix.lower()


def convert(src: Path, dst: Path) -> None:
    src, dst = Path(src), Path(dst)
    if not src.exists(): raise FileNotFoundError(f"找不到输入文件：{src}")
    if src.is_dir(): raise ValueError("暂不支持直接转换文件夹，请先压缩成 ZIP/TAR。")
    if src.resolve() == dst.resolve(): raise ValueError("输入文件和输出文件不能相同。")
    dst.parent.mkdir(parents=True, exist_ok=True)
    s, d = format_name(src), format_name(dst)
    if s == d:
        shutil.copy2(src, dst); return
    if s == ".pdf" and d == ".docx":
        pdf_to_docx(src, dst); return
    if s in IMAGES and (d in IMAGES or d == ".pdf"):
        image_convert(src, dst); return
    if s == ".pdf" and d in IMAGES:
        pdf_to_image(src, dst); return
    if s in TEXT and d in TEXT:
        text_convert(src, dst); return
    if (s in OFFICE or d in OFFICE) and d not in MEDIA:
        office_convert(src, dst); return
    if s in MEDIA or d in MEDIA:
        ffmpeg_convert(src, dst); return
    if s in ARCHIVES or d in ARCHIVES:
        archive_convert(src, dst); return
    raise ValueError(f"暂不支持：{s} → {d}")


def image_convert(src: Path, dst: Path) -> None:
    with Image.open(src) as original:
        im = original.convert("RGB") if dst.suffix.lower() in {".jpg", ".jpeg"} and original.mode in {"RGBA", "LA", "P"} else original
        save_kwargs = {"quality": 95} if dst.suffix.lower() in {".jpg", ".jpeg", ".webp"} else {}
        im.save(dst, **save_kwargs)


def pdf_to_image(src: Path, dst: Path) -> None:
    doc = fitz.open(src)
    try:
        if len(doc) == 0: raise ValueError("PDF 没有页面")
        doc[0].get_pixmap(dpi=150, alpha=False).save(dst)
    finally:
        doc.close()


def _pdf_font_name(span: dict) -> str:
    name = span.get("font", "Microsoft YaHei")
    if "+" in name:
        name = name.split("+", 1)[-1]
    mapping = {
        "HarmonyOS_Sans_SC": "Microsoft YaHei",
        "NotoSerifCJKjp-Regular": "SimSun",
        "NotoSerifCJKsc-Regular": "SimSun",
        "SourceHanSansCN-Regular": "Microsoft YaHei",
    }
    return mapping.get(name, name or "Microsoft YaHei")


def _pdf_span_run_xml(span: dict) -> str:
    text = html.escape(span.get("text", ""), quote=False)
    if not text:
        return ""
    size = float(span.get("size", 10) or 10)
    size = max(5, min(size, 72))
    font = html.escape(_pdf_font_name(span), quote=True)
    flags = int(span.get("flags", 0))
    # char_flags is not a reliable bold/italic indicator across embedded CJK fonts.
    bold = bool(flags & 16)
    italic = bool(flags & 2)
    props = (
        f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:eastAsia="{font}"/>'
        f'<w:sz w:val="{round(size * 2)}"/><w:szCs w:val="{round(size * 2)}"/>'
        + ("<w:b/>" if bold else "")
        + ("<w:i/>" if italic else "")
    )
    return f'<w:r><w:rPr>{props}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'


def _add_pdf_textbox(paragraph, block: dict, page_width_pt: float, shape_id: int) -> None:
    """Add an absolutely positioned, editable Word text box for one PDF text block."""
    x0, y0, x1, y1 = (float(v) for v in block["bbox"])
    width = max(8.0, x1 - x0 + 2.0)
    height = max(8.0, y1 - y0 + 4.0)
    line_xml = []
    for line_index, line in enumerate(block.get("lines", [])):
        if line_index:
            line_xml.append("<w:br/>")
        for span in line.get("spans", []):
            line_xml.append(_pdf_span_run_xml(span))
    if not line_xml:
        return

    center = (x0 + x1) / 2.0
    align = "center" if abs(center - page_width_pt / 2.0) < page_width_pt * 0.08 else "left"
    content = "".join(line_xml)
    xml = f'''<w:pict {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml">
      <v:shape id="pdfText{shape_id}" type="#_x0000_t202"
        style="position:absolute;left:{x0:.2f}pt;top:{y0:.2f}pt;width:{width:.2f}pt;height:{height:.2f}pt;z-index:2;mso-wrap-style:none"
        stroked="f" filled="f">
        <v:textbox inset="0pt,0pt,0pt,0pt">
          <w:txbxContent>
            <w:p>
              <w:pPr><w:jc w:val="{align}"/><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>
              {content}
            </w:p>
          </w:txbxContent>
        </v:textbox>
      </v:shape>
    </w:pict>'''
    paragraph.add_run()._r.append(parse_xml(xml))


def _redacted_page_png(page: fitz.Page, text_blocks: list[dict], dpi: int = 150) -> bytes:
    """Render the page after removing only PDF text, keeping lines, drawings and images."""
    work = fitz.open()
    try:
        new_page = work.new_page(width=page.rect.width, height=page.rect.height)
        new_page.show_pdf_page(new_page.rect, page.parent, page.number)
        for block in text_blocks:
            rect = fitz.Rect(block["bbox"])
            rect.x0 -= 0.7; rect.y0 -= 0.7; rect.x1 += 0.7; rect.y1 += 0.7
            new_page.add_redact_annot(rect, fill=None)
        if text_blocks:
            new_page.apply_redactions(images=0, graphics=0, text=0)
        pix = new_page.get_pixmap(dpi=dpi, alpha=False)
        return pix.tobytes("png")
    finally:
        work.close()


def pdf_to_docx(src: Path, dst: Path) -> None:
    """Convert PDF to a layout-preserving, editable DOCX.

    Each PDF page becomes a fixed-size Word section. The original page is used
    as a background after PDF text is redacted, preserving drawings, rules,
    check marks and signatures. Extracted text is then placed back as editable
    absolute-positioned Word text boxes. This is substantially more faithful
    for forms and scanned/hybrid PDFs than normal paragraph-flow conversion.
    """
    pdf = fitz.open(src)
    try:
        if len(pdf) == 0:
            raise ValueError("PDF 没有页面")
        docx = Document()
        for page_index, page in enumerate(pdf):
            if page_index:
                section = docx.add_section(WD_SECTION.NEW_PAGE)
            else:
                section = docx.sections[0]
            rect = page.rect
            section.page_width = Inches(rect.width / 72)
            section.page_height = Inches(rect.height / 72)
            section.top_margin = Inches(0)
            section.bottom_margin = Inches(0)
            section.left_margin = Inches(0)
            section.right_margin = Inches(0)
            section.header_distance = Inches(0)
            section.footer_distance = Inches(0)

            blocks = page.get_text("dict").get("blocks", [])
            text_blocks = [
                b for b in blocks
                if b.get("type") == 0
                and any(s.get("text", "").strip() for l in b.get("lines", []) for s in l.get("spans", []))
            ]

            bg_data = _redacted_page_png(page, text_blocks, dpi=150)
            paragraph = docx.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 0.01

            image_run = paragraph.add_run()
            image_run.add_picture(BytesIO(bg_data), width=Inches(rect.width / 72))
            rid = image_run._r.xpath('.//a:blip/@r:embed')[0]
            image_run._r.getparent().remove(image_run._r)
            bg_xml = f'''<w:pict {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
              <v:shape id="pdfPage{page_index}" type="#_x0000_t75"
                style="position:absolute;left:0pt;top:0pt;width:{rect.width:.2f}pt;height:{rect.height:.2f}pt;z-index:-1;mso-wrap-style:none"
                stroked="f" filled="f">
                <v:imagedata r:id="{rid}" o:title="PDF page"/>
              </v:shape>
            </w:pict>'''
            paragraph.add_run()._r.append(parse_xml(bg_xml))

            for shape_id, block in enumerate(text_blocks, start=1):
                _add_pdf_textbox(paragraph, block, rect.width, shape_id)

        docx.save(dst)
    finally:
        pdf.close()


def _read_text_data(src: Path):
    raw = src.read_text(encoding="utf-8-sig")
    suffix = src.suffix.lower()
    if suffix == ".json": return json.loads(raw)
    if suffix in {".yaml", ".yml"}: return yaml.safe_load(raw)
    if suffix == ".csv": return list(csv.DictReader(raw.splitlines()))
    if suffix == ".xml": return xml_to_dict(ET.fromstring(raw))
    return raw


def text_convert(src: Path, dst: Path) -> None:
    obj = _read_text_data(src); suffix = dst.suffix.lower()
    if suffix == ".json": dst.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    elif suffix in {".yaml", ".yml"}: dst.write_text(yaml.safe_dump(obj, allow_unicode=True, sort_keys=False), encoding="utf-8")
    elif suffix == ".csv":
        if not isinstance(obj, list) or not all(isinstance(row, dict) for row in obj): raise ValueError("只有对象数组/表格数据才能可靠转换为 CSV。")
        keys = list(dict.fromkeys(key for row in obj for key in row))
        with dst.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=keys, extrasaction="ignore"); writer.writeheader(); writer.writerows(obj)
    elif suffix == ".xml": dst.write_text(dict_to_xml(obj), encoding="utf-8")
    else:
        if not isinstance(obj, str): obj = json.dumps(obj, ensure_ascii=False, indent=2)
        dst.write_text(obj, encoding="utf-8")


def xml_to_dict(elem: ET.Element):
    children = list(elem)
    if not children: return elem.text or ""
    out = {}
    for child in children:
        value = xml_to_dict(child)
        if child.tag in out:
            if not isinstance(out[child.tag], list): out[child.tag] = [out[child.tag]]
            out[child.tag].append(value)
        else: out[child.tag] = value
    return {elem.tag: out}


def dict_to_xml(obj) -> str:
    if not isinstance(obj, dict) or len(obj) != 1: raise ValueError("XML 输出需要一个唯一的根节点对象。")
    root_name, value = next(iter(obj.items())); root = ET.Element(str(root_name)); _fill_xml(root, value)
    return ET.tostring(root, encoding="unicode")


def _fill_xml(parent: ET.Element, value) -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            values = child if isinstance(child, list) else [child]
            for item in values: node = ET.SubElement(parent, str(key)); _fill_xml(node, item)
    elif isinstance(value, list):
        for item in value: node = ET.SubElement(parent, "item"); _fill_xml(node, item)
    elif value is not None: parent.text = str(value)


def office_convert(src: Path, dst: Path) -> None:
    soffice = find_soffice()
    if not soffice: raise RuntimeError("未检测到 LibreOffice。请在“Office 依赖”中自动安装，或手动安装后重试。")
    with tempfile.TemporaryDirectory() as td:
        cmd = [soffice, "--headless", "--convert-to", dst.suffix.lstrip("."), "--outdir", td, str(src)]
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        produced = Path(td) / f"{src.stem}{dst.suffix}"
        if p.returncode != 0 or not produced.exists(): raise RuntimeError((p.stderr or p.stdout or "LibreOffice 转换失败").strip())
        shutil.copy2(produced, dst)


def ffmpeg_convert(src: Path, dst: Path) -> None:
    ffmpeg = shutil.which("ffmpeg")
    if not ffmpeg: raise RuntimeError("音视频转换需要 FFmpeg，并确保 ffmpeg 在 PATH 中。")
    p = subprocess.run([ffmpeg, "-hide_banner", "-y", "-i", str(src), str(dst)], capture_output=True, text=True, timeout=1800)
    if p.returncode != 0: raise RuntimeError((p.stderr or p.stdout or "FFmpeg 转换失败")[-4000:])


def _extract_archive(src: Path, workdir: Path) -> None:
    suffix = format_name(src)
    if suffix == ".zip":
        with zipfile.ZipFile(src) as z: z.extractall(workdir)
    elif suffix in {".tar", ".tar.gz", ".tgz", ".tar.bz2", ".tar.xz"}:
        with tarfile.open(src, "r:*") as tar:
            if hasattr(tarfile, "data_filter"): tar.extractall(workdir, filter="data")
            else: tar.extractall(workdir)
    else: raise ValueError(f"暂不支持解压：{suffix}")


def _archive_directory(workdir: Path, dst: Path) -> None:
    suffix = format_name(dst)
    if suffix == ".zip":
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
            for path in workdir.rglob("*"):
                if path.is_file(): z.write(path, path.relative_to(workdir))
        return
    mode = {".tar": "w", ".tar.gz": "w:gz", ".tgz": "w:gz", ".tar.bz2": "w:bz2", ".tar.xz": "w:xz"}.get(suffix)
    if not mode: raise ValueError(f"暂不支持生成归档：{suffix}")
    with tarfile.open(dst, mode) as tar:
        for path in workdir.rglob("*"): tar.add(path, arcname=path.relative_to(workdir))


def archive_convert(src: Path, dst: Path) -> None:
    supported = {".zip", ".tar", ".tar.gz", ".tgz", ".tar.bz2", ".tar.xz"}
    if format_name(src) not in supported: raise ValueError("输入压缩格式目前仅支持 ZIP/TAR/TAR.GZ/TGZ/TAR.BZ2/TAR.XZ。")
    if format_name(dst) not in supported: raise ValueError("输出压缩格式目前仅支持 ZIP/TAR/TAR.GZ/TGZ/TAR.BZ2/TAR.XZ。")
    with tempfile.TemporaryDirectory() as td:
        workdir = Path(td) / "payload"; workdir.mkdir(); _extract_archive(src, workdir); _archive_directory(workdir, dst)
