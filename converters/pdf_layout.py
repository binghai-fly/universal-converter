from __future__ import annotations

import html
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt


def _pdf_font_name(span: dict) -> str:
    name = span.get("font", "Microsoft YaHei")
    if "+" in name:
        name = name.split("+", 1)[-1]
    return {
        "HarmonyOS_Sans_SC": "Microsoft YaHei",
        "NotoSerifCJKjp-Regular": "SimSun",
        "NotoSerifCJKsc-Regular": "SimSun",
        "SourceHanSansCN-Regular": "Microsoft YaHei",
        "SimSun,宋体": "SimSun",
        "SimHei,黑体": "SimHei",
    }.get(name, name or "Microsoft YaHei")


def _pdf_color(span: dict) -> str | None:
    value = span.get("color")
    if isinstance(value, int):
        return f"{(value >> 16) & 255:02X}{(value >> 8) & 255:02X}{value & 255:02X}"
    return None


def _pdf_span_run_xml(span: dict) -> str:
    text = html.escape(span.get("text", ""), quote=False)
    if not text:
        return ""
    size = max(5.0, min(float(span.get("size", 10) or 10), 72.0))
    font = html.escape(_pdf_font_name(span), quote=True)
    flags = int(span.get("flags", 0))
    char_flags = int(span.get("char_flags", 0))
    props = (
        f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" '
        f'w:eastAsia="{font}" w:cs="{font}"/>'
        f'<w:sz w:val="{round(size * 2)}"/>'
        f'<w:szCs w:val="{round(size * 2)}"/>'
    )
    if flags & 16:
        props += "<w:b/><w:bCs/>"
    if flags & 2:
        props += "<w:i/><w:iCs/>"
    color = _pdf_color(span)
    if color and color != "000000":
        props += f'<w:color w:val="{color}"/>'
    if char_flags & 4:
        props += '<w:u w:val="single"/>'
    return f'<w:r><w:rPr>{props}</w:rPr><w:t xml:space="preserve">{text}</w:t></w:r>'


def _add_pdf_textbox(paragraph, line: dict, shape_id: int, page_width_pt: float) -> None:
    spans = line.get("spans", [])
    if not spans:
        return
    content = "".join(_pdf_span_run_xml(span) for span in spans)
    if not content:
        return

    x0, y0, x1, y1 = (float(v) for v in line["bbox"])
    width = max(5.0, x1 - x0 + 3.0)
    height = max(14.0, y1 - y0 + 8.0)
    center = (x0 + x1) / 2.0
    align = "center" if abs(center - page_width_pt / 2.0) < page_width_pt * 0.08 else "left"

    xml = f'''<w:pict {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml">
      <v:shape id="pdfText{shape_id}" type="#_x0000_t202"
        style="position:absolute;margin-left:{x0 - 1:.2f}pt;margin-top:{y0 - 1:.2f}pt;
        width:{width:.2f}pt;height:{height:.2f}pt;z-index:2;mso-wrap-style:none;
        mso-position-horizontal-relative:page;mso-position-vertical-relative:page"
        stroked="f" filled="f">
        <v:textbox style="mso-fit-shape-to-text:t;mso-margin-left:0;mso-margin-right:0;
        mso-margin-top:0;mso-margin-bottom:0">
          <w:txbxContent>
            <w:p>
              <w:pPr>
                <w:jc w:val="{align}"/>
                <w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>
              </w:pPr>
              {content}
            </w:p>
          </w:txbxContent>
        </v:textbox>
      </v:shape>
    </w:pict>'''
    paragraph.add_run()._r.append(parse_xml(xml))


def _redacted_page_png(page: fitz.Page, line_rects: list[tuple[float, float, float, float]], dpi: int = 150) -> bytes:
    work = fitz.open()
    try:
        new_page = work.new_page(width=page.rect.width, height=page.rect.height)
        new_page.show_pdf_page(new_page.rect, page.parent, page.number)
        for raw in line_rects:
            rect = fitz.Rect(raw)
            rect.x0 -= 0.6
            rect.y0 -= 0.6
            rect.x1 += 0.6
            rect.y1 += 0.6
            new_page.add_redact_annot(rect, fill=None)
        if line_rects:
            new_page.apply_redactions(images=0, graphics=0, text=0)
        pix = new_page.get_pixmap(dpi=dpi, alpha=False)
        return pix.tobytes("png")
    finally:
        work.close()


def convert_pdf_to_docx(src: Path, dst: Path) -> None:
    """Create a visually stable DOCX with editable text overlays.

    The original PDF page is preserved as a high-resolution background after
    text removal. Extracted PDF text is reconstructed as editable Word text
    boxes at the original page coordinates, preserving fixed-form layouts.
    """
    pdf = fitz.open(src)
    try:
        if not len(pdf):
            raise ValueError("PDF 没有页面")

        docx = Document()
        for page_index, page in enumerate(pdf):
            section = docx.sections[0] if page_index == 0 else docx.add_section(WD_SECTION.NEW_PAGE)
            rect = page.rect
            section.page_width = Inches(rect.width / 72)
            section.page_height = Inches(rect.height / 72)
            section.top_margin = Inches(0)
            section.bottom_margin = Inches(0)
            section.left_margin = Inches(0)
            section.right_margin = Inches(0)
            section.header_distance = Inches(0)
            section.footer_distance = Inches(0)

            blocks = page.get_text("dict").get("blocks", [])
            lines = []
            for block in blocks:
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    if any(span.get("text", "").strip() for span in line.get("spans", [])):
                        lines.append(line)

            background = _redacted_page_png(page, [tuple(line["bbox"]) for line in lines])
            paragraph = docx.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 0.01

            image_run = paragraph.add_run()
            image_run.add_picture(BytesIO(background), width=Inches(rect.width / 72))
            rid = image_run._r.xpath('.//a:blip/@r:embed')[0]
            image_run._r.getparent().remove(image_run._r)

            background_xml = f'''<w:pict {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"
              xmlns:o="urn:schemas-microsoft-com:office:office"
              xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
              <v:shape id="pdfPage{page_index}" type="#_x0000_t75"
                style="position:absolute;margin-left:0pt;margin-top:0pt;width:{rect.width:.2f}pt;
                height:{rect.height:.2f}pt;z-index:-1;mso-wrap-style:none;
                mso-position-horizontal-relative:page;mso-position-vertical-relative:page"
                stroked="f" filled="f">
                <v:imagedata r:id="{rid}" o:title="PDF page"/>
              </v:shape>
            </w:pict>'''
            paragraph.add_run()._r.append(parse_xml(background_xml))

            for shape_id, line in enumerate(lines, start=1 + page_index * 10000):
                _add_pdf_textbox(paragraph, line, shape_id, rect.width)

        docx.save(dst)
    finally:
        pdf.close()
